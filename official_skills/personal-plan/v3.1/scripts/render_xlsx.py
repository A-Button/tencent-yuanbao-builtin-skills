#!/usr/bin/env python3
"""把计划 JSON 写成带基础样式的 XLSX。内容由调用方提供，本脚本只负责写表。

用法（从本 Skill 目录，或写出绝对路径）：
    python3 scripts/render_xlsx.py <数据.json> <中文计划名>.xlsx

写出成功后打印 5 行结构报告（写出路径、表名、执行表、填写列、预填）。
结构项只看这 5 行，不要另写检查脚本。有 FAIL 只改 JSON 再跑一次，然后停。
不要读或改本脚本。

JSON 顶层：
{
  "sheets": [
    {
      "name": "计划概览",
      "blocks": [
        {"type": "heading", "text": "目标与判断"},
        {"type": "kv", "rows": [["目标", "..."], ["期限", "..."]]},
        {"type": "heading", "text": "阶段路线"},
        {"type": "table", "headers": ["阶段", "日期范围", "重点", "通过标准"], "rows": [["基础", "第1–2周", "...", "..."]]}
      ]
    },
    {
      "name": "执行计划",
      "blocks": [
        {"type": "note", "text": "前两周按天，之后按周。"},
        {"type": "table", "headers": ["日期", "科目/模块", "学习内容", "题量或产出", "耗时", "完成标准", "状态", "实际结果", "反馈/问题"], "rows": [["2026-08-24", "数学", "有理数运算专项", "60题+错题本建册", "1.5h", "连续20题无错", "", "", ""]]}
      ]
    }
  ]
}

block.type：
- heading：区块标题，字段 text
- kv：标签+内容，字段 rows 为 [标签, 内容]
- table：明细表，字段 headers、rows
- note：整行说明，字段 text

执行表必须是带日期或周次、且有数据行的 table。列头按领域写（学习用学习内容 / 题量或产出，健身用动作 / 组次），报告不要求必须叫「动作」或「数量或输出」。
填写列「状态」「实际结果」「反馈/问题」整本至少出现在一张 table 里，格子留空。不要预填「未开始」。标明「示例」的示范行可以带字，不算预填。
"""

from __future__ import annotations

import json
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_NAME = "微软雅黑"
TITLE_COLOR = "1F3864"
HEADER_BG = "2F5597"
HEADER_FG = "FFFFFF"
LABEL_BG = "D9E2F3"
RECORD_BG = "F2F2F2"
TEXT = "1F1F1F"

RECORD_HEADERS = {"状态", "实际结果", "反馈/问题", "体重", "围度"}
FILL_HEADERS = ("状态", "实际结果", "反馈/问题")
CALENDAR_HEADERS = frozenset({
    "日期", "周次", "星期", "训练日", "日期/周次", "日期或周次", "时段",
})
RANGE_HEADERS = frozenset({"日期范围", "时间范围", "时间"})
TASK_HINTS = frozenset({
    "学习内容", "科目", "科目/模块", "科目或模块", "知识点",
    "动作", "训练任务", "任务内容", "任务", "训练内容", "本周重点",
    "岗位", "投递动作", "对象", "内容",
})
SHORT_HEADERS = {
    "日期", "星期", "周次", "时段", "时长", "耗时", "RPE", "优先级", "类型",
    "状态", "序号", "阶段", "频率", "轮次", "检查时点", "科目", "科目/模块",
}
LONG_HEADERS = {
    "内容", "说明", "说明/依据", "任务内容", "训练任务", "完成标准", "核心策略",
    "通过标准", "实际结果", "反馈/问题", "注意事项", "阶段目标",
    "学习内容", "题量或产出",
}

WRAP_HEADERS = LONG_HEADERS | {"动作", "数量或输出", "最小方案", "依赖", "投递动作"}


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _sheet_name(value: str) -> str:
    cleaned = "".join(ch for ch in value.strip() if ch not in r"\/*?:[]") or "计划"
    return cleaned[:31]


def _visual_len(value: object) -> int:
    text = _text(value)
    widths = []
    for line in text.splitlines() or [""]:
        widths.append(sum(2 if unicodedata.east_asian_width(ch) in "WFA" else 1 for ch in line))
    return max(widths, default=0)


def _estimate_width(header: str, values: list[object]) -> float:
    name = header.strip()
    longest = max([_visual_len(name), *(_visual_len(v) for v in values)], default=8)
    if name in SHORT_HEADERS:
        return min(16, max(10, longest + 2))
    if name in LONG_HEADERS or name in WRAP_HEADERS:
        return min(42, max(28, longest * 0.7 + 4))
    return min(30, max(14, longest * 0.7 + 4))


def _font(size: int = 11, bold: bool = False, color: str = TEXT) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _align(wrap: bool = False, center: bool = False) -> Alignment:
    return Alignment(
        horizontal="center" if center else "left",
        vertical="top",
        wrap_text=wrap,
    )


def _apply(cell, *, font: Font, fill: PatternFill | None = None, align: Alignment | None = None) -> None:
    cell.font = font
    if fill is not None:
        cell.fill = fill
    cell.alignment = align or _align()


def _max_cols(blocks: list[dict]) -> int:
    widest = 2
    for block in blocks:
        kind = block.get("type")
        if kind == "table":
            widest = max(widest, len(block.get("headers") or []))
        elif kind == "kv":
            widest = max(widest, 2)
    return widest


def _style_header_row(ws, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        _apply(
            cell,
            font=_font(11, bold=True, color=HEADER_FG),
            fill=_fill(HEADER_BG),
            align=_align(center=True),
        )


def _write_heading(ws, row: int, text: str, cols: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    _apply(cell, font=_font(14, bold=True, color=TITLE_COLOR), fill=_fill(LABEL_BG))
    for col in range(2, cols + 1):
        _apply(ws.cell(row, col), font=_font(14, bold=True, color=TITLE_COLOR), fill=_fill(LABEL_BG))
    return row + 1


def _write_note(ws, row: int, text: str, cols: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    _apply(cell, font=_font(), align=_align(wrap=True))
    return row + 1


def _write_kv(ws, row: int, rows: list, cols: int) -> int:
    for item in rows:
        if not isinstance(item, (list, tuple)) or len(item) < 2:
            raise ValueError("kv.rows 每一项必须是 [标签, 内容]")
        label, value = item[0], item[1]
        label_cell = ws.cell(row, 1, _text(label))
        _apply(label_cell, font=_font(bold=True), fill=_fill(LABEL_BG))
        if cols > 2:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=cols)
        value_cell = ws.cell(row, 2, _text(value))
        _apply(value_cell, font=_font(), align=_align(wrap=True))
        for col in range(3, cols + 1):
            _apply(ws.cell(row, col), font=_font())
        row += 1
    return row


def _write_table(ws, row: int, headers: list[str], rows: list, cols: int) -> int:
    if not headers:
        raise ValueError("table.headers 不能为空")
    _style_header_row(ws, row, headers)
    for col in range(len(headers) + 1, cols + 1):
        _apply(ws.cell(row, col), font=_font(11, bold=True, color=HEADER_FG), fill=_fill(HEADER_BG))
    row += 1
    record_cols = {idx for idx, name in enumerate(headers, 1) if name in RECORD_HEADERS}
    wrap_cols = {idx for idx, name in enumerate(headers, 1) if name in WRAP_HEADERS}
    for item in rows:
        values = list(item) if isinstance(item, (list, tuple)) else [_text(item)]
        if len(values) < len(headers):
            values.extend([""] * (len(headers) - len(values)))
        for col, header in enumerate(headers, 1):
            raw = values[col - 1]
            text = _text(raw)
            cell = ws.cell(row, col, text)
            is_record = col in record_cols
            _apply(
                cell,
                font=_font(),
                fill=_fill(RECORD_BG) if is_record else None,
                align=_align(wrap=col in wrap_cols),
            )
        row += 1
    return row


def _apply_widths(ws, headers_by_col: list[list[str]], values_by_col: list[list[object]]) -> None:
    for idx, (headers, values) in enumerate(zip(headers_by_col, values_by_col), 1):
        header = headers[0] if headers else ""
        # 同一列出现过短标签和长说明时，按更长的那头估，避免下面多列表被挤
        if any(h in LONG_HEADERS or h in WRAP_HEADERS for h in headers):
            header = next(h for h in headers if h in LONG_HEADERS or h in WRAP_HEADERS)
        elif any(h in SHORT_HEADERS for h in headers) and not any(
            h not in SHORT_HEADERS and h for h in headers
        ):
            header = next(h for h in headers if h in SHORT_HEADERS)
        ws.column_dimensions[get_column_letter(idx)].width = _estimate_width(header, values)


def _norm_header(value: str) -> str:
    return value.replace(" ", "").replace("／", "/")


def _header_kind(header: str) -> str | None:
    name = _norm_header(header)
    if name in {_norm_header(item) for item in RANGE_HEADERS} or ("日期" in name and "范围" in name):
        return "range"
    if name in {_norm_header(item) for item in CALENDAR_HEADERS}:
        return "calendar"
    if name.startswith("日期") or name.startswith("周次"):
        return "calendar"
    return None


def _is_example_row(row: list | tuple) -> bool:
    texts = [_text(cell).strip() for cell in row]
    if not any(texts):
        return True
    for text in texts:
        if text.startswith("示例") or text.startswith("例：") or text.startswith("例:") or text.startswith("示范"):
            return True
    return False


def _is_exec_table(headers: list[str], rows: list) -> bool:
    header_set = {_norm_header(name) for name in headers if name}
    kinds = {_header_kind(name) for name in headers}
    if "calendar" not in kinds and "range" not in kinds:
        return False
    data_rows = [row for row in rows if not _is_example_row(row)]
    if not data_rows:
        return False
    # 概览里的阶段路线：有「阶段」、没有任务列，不算执行表
    if "阶段" in header_set and not (header_set & {_norm_header(item) for item in TASK_HINTS}):
        return False
    return True


def _structure_report(data: dict) -> list[str]:
    sheets = data.get("sheets") if isinstance(data.get("sheets"), list) else []
    names: list[str] = []
    exec_rows = 0
    fill_found = False
    fill_prefilled = False
    kv_empty = False

    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue
        names.append(_sheet_name(_text(sheet.get("name") or "计划")))
        for block in sheet.get("blocks") or []:
            if not isinstance(block, dict):
                continue
            kind = block.get("type")
            if kind == "kv":
                for item in block.get("rows") or []:
                    if isinstance(item, (list, tuple)) and len(item) >= 2 and not _text(item[1]).strip():
                        kv_empty = True
            elif kind == "table":
                headers = [_text(h).strip() for h in (block.get("headers") or [])]
                header_set = set(headers)
                rows = [r for r in (block.get("rows") or []) if isinstance(r, (list, tuple))]
                if _is_exec_table(headers, rows):
                    exec_rows += len([r for r in rows if not _is_example_row(r)])
                if all(name in header_set for name in FILL_HEADERS):
                    fill_found = True
                    positions = {name: headers.index(name) for name in FILL_HEADERS}
                    for row in rows:
                        if _is_example_row(row):
                            continue
                        for name in FILL_HEADERS:
                            pos = positions[name]
                            if pos < len(row) and _text(row[pos]).strip():
                                fill_prefilled = True

    exec_line = (
        "执行表: OK"
        if exec_rows > 0
        else "执行表: FAIL 没有带日期或周次、且有数据行的排期表"
    )
    fill_line = (
        "填写列: OK"
        if fill_found
        else "填写列: FAIL 整本没有同时含「状态」「实际结果」「反馈/问题」的表"
    )
    extras = []
    if fill_prefilled:
        extras.append("填写列有字（「未开始」也不行）")
    if kv_empty:
        extras.append("概览有内容列空着")
    prefill_line = "预填: FAIL " + "；".join(extras) if extras else "预填: OK"
    return [
        f"表: {', '.join(names) if names else '（无）'}",
        exec_line,
        fill_line,
        prefill_line,
    ]


def render(data: dict, output: Path) -> Path:
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("JSON 顶层必须有非空的 sheets")

    wb = Workbook()
    wb.remove(wb.active)

    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValueError("每个 sheet 必须是对象")
        name = _sheet_name(_text(sheet.get("name") or "计划"))
        blocks = sheet.get("blocks")
        if not isinstance(blocks, list) or not blocks:
            raise ValueError(f"{name}: blocks 不能为空")

        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        cols = _max_cols(blocks)
        row = 1
        headers_by_col: list[list[str]] = [[] for _ in range(cols)]
        values_by_col: list[list[object]] = [[] for _ in range(cols)]
        table_count = 0

        for block in blocks:
            if not isinstance(block, dict):
                raise ValueError(f"{name}: 每个 block 必须是对象")
            kind = block.get("type")
            if kind == "heading":
                row = _write_heading(ws, row, _text(block.get("text")), cols)
            elif kind == "note":
                row = _write_note(ws, row, _text(block.get("text")), cols)
            elif kind == "kv":
                rows = block.get("rows") or []
                row = _write_kv(ws, row, rows, cols)
                headers_by_col[0].append("标签")
                headers_by_col[1].append("内容")
                for item in rows:
                    values_by_col[0].append(item[0])
                    values_by_col[1].append(item[1])
            elif kind == "table":
                headers = [ _text(h) for h in (block.get("headers") or []) ]
                rows = block.get("rows") or []
                row = _write_table(ws, row, headers, rows, cols)
                table_count += 1
                for idx, header in enumerate(headers):
                    headers_by_col[idx].append(header)
                    values_by_col[idx].append(header)
                    for item in rows:
                        if isinstance(item, (list, tuple)) and idx < len(item):
                            values_by_col[idx].append(item[idx])
            else:
                raise ValueError(f"{name}: 不支持的 block.type: {kind}")
            row += 1

        _apply_widths(ws, headers_by_col, values_by_col)
        if table_count == 1 and not any(b.get("type") == "kv" for b in blocks):
            # 整张就是一张明细表时冻首个表头行
            for probe in range(1, ws.max_row + 1):
                if ws.cell(probe, 1).fill.fgColor and str(ws.cell(probe, 1).fill.fgColor.rgb or "").endswith(HEADER_BG):
                    ws.freeze_panes = f"A{probe + 1}"
                    break

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    check = load_workbook(output, read_only=True)
    names = check.sheetnames
    check.close()
    if not names:
        raise ValueError("写出的工作簿没有工作表")
    return output


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("用法: python3 scripts/render_xlsx.py <数据.json> <输出.xlsx>", file=sys.stderr)
        return 2
    src = Path(argv[1])
    dest = Path(argv[2])
    if not src.exists():
        print(f"找不到数据文件: {src}", file=sys.stderr)
        return 1
    try:
        data = json.loads(src.read_text(encoding="utf-8"))
        path = render(data, dest)
    except (json.JSONDecodeError, ValueError, KeyError, TypeError) as exc:
        print(f"渲染失败: {exc}", file=sys.stderr)
        return 1
    print(f"已写出 {path}")
    for line in _structure_report(data):
        print(line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
